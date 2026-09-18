"""Pull specifications out of the supplier's product names.

The supplier writes every parameter into the name itself, in a form that is
repeated across the whole catalogue:

    Фитинг тройник Y-образный, 8мм x 8мм x 8мм, T=(-20...+80)°C,
    P=(-1...15) бар, корпус - никел.латунь, цанга - никел.латунь

so each parameter can be read back out and put in its own row.
"""

import re

RANGE = r'[-+]?\d+(?:[.,]\d+)?\s*\.\.\.\s*[-+]?\d+(?:[.,]\d+)?'
X = '[xх×]'  # the file mixes the Latin x with the Cyrillic х
M = '[MМ]'   # and the Latin M with the Cyrillic М


def tidy(value):
    return re.sub(r'\s+', ' ', value).strip(' ,;.').strip()


def latin(value):
    return value.replace('х', 'x').replace('×', 'x').replace('М', 'M')


def temperature(name):
    found = re.search(rf'T\s*=\s*\(?\s*({RANGE})\s*\)?\s*°?\s*C', name, re.IGNORECASE)
    if found:
        return f'{tidy(found.group(1))} °C'

    found = re.search(r'от\s*([-+]?\d+)\s*°?\s*C?\s*до\s*([-+]?\d+)\s*°?\s*C', name, re.IGNORECASE)
    if found:
        low, high = found.group(1), found.group(2)
        high = high if high.startswith(('-', '+')) else f'+{high}'
        return f'{low}...{high} °C'

    found = re.search(rf'(?<![\w=])({RANGE})\s*°C', name)
    return f'{tidy(found.group(1))} °C' if found else None


def pressure(name, prefix=''):
    found = re.search(
        rf'P{prefix}\s*=\s*\(?\s*({RANGE})\s*\)?\s*бар', name, re.IGNORECASE)
    if found:
        return f'{tidy(found.group(1))} бар'
    if not prefix:
        found = re.search(rf'(?<![\w=])({RANGE})\s*бар', name, re.IGNORECASE)
        if found:
            return f'{tidy(found.group(1))} бар'
    return None


def material(name, part):
    found = re.search(rf'{part}\s*[-–—]\s*([^,;()]+)', name, re.IGNORECASE)
    return tidy(found.group(1)) if found else None


def voltage(name):
    patterns = [
        rf'([=~]?\s*{RANGE}\s*(?:В|VAC|VDC|Вольт))(?![\w])',
        r'([=~]?\s*\d+\s*[-–~]\s*\d+\s*(?:В|VAC|VDC|Вольт))(?![\w])',
        r'([=~]?\s*\d+(?:[.,]\d+)?\s*(?:В|VAC|VDC|Вольт))(?![\w])',
    ]
    for pattern in patterns:
        found = re.search(pattern, name, re.IGNORECASE)
        if found:
            value = tidy(found.group(1))
            # «=24 В» direct current, «~230 В» alternating — keep the sign,
            # put one space before the unit
            return re.sub(r'\s*(В|VAC|VDC|Вольт)$', r' \1',
                          re.sub(r'^([=~])\s*', r'\1', value))
    return None


INCH = r'(?:\'\'|"|”)'
FRACTION = r'\d+/\d+'


def quotes(value):
    return tidy(value).replace("''", '"').replace('”', '"')


def thread(name):
    patterns = [
        rf'\b\d+\s+{FRACTION}\s*{INCH}',            # 1 1/4"
        rf'\b[GRК]\s*\d*\s*{FRACTION}\s*{INCH}?',   # G 1/4", R 3/8
        rf'\b{M}\d+(?:{X}\d+(?:[.,]\d+)?)?',        # M20x1.5, М8х1
        rf'(?<![\d/]){FRACTION}\s*{INCH}',          # 1/2"
        rf'(?<![\d/]){FRACTION}\s*(?=\(?\s*(?:внутр|наруж))',  # 1/2 (наруж.резьба)
    ]
    for pattern in patterns:
        found = re.search(pattern, name)
        if found:
            return latin(quotes(found.group(0)))
    return None


CONNECTION = r'\d+(?:-\d+)?\s*мм'


def size(name):
    """The dimension the name carries, with the label that fits its shape."""
    found = re.search(rf'{CONNECTION}(?:\s*{X}\s*{CONNECTION})+', name)
    if found:
        return 'Присоединение', latin(tidy(found.group(0)))

    # «8мм x 3/8 (наруж.резьба)» — millimetres on one side, a thread on the other
    found = re.search(rf'({CONNECTION})\s*{X}\s*(?=[GRК]?\s*{FRACTION})', name)
    if found:
        return 'Присоединение', latin(tidy(found.group(1)))

    found = re.search(r'\bDN\s*\d+', name)
    if found:
        return 'Условный проход', tidy(found.group(0))

    for pattern in (r'для\s+пневмоцилиндров\s*D\s*=\s*([\d\s-]+)\s*(?:мм)?',
                    r'D\s*=\s*([\d\s-]+?)\s*мм\s*,?\s*для\s+пневмоцилиндров'):
        found = re.search(pattern, name, re.IGNORECASE)
        if found:
            return 'Для цилиндров', f'D = {tidy(found.group(1))} мм'

    found = re.search(r'на\s+(?:пневматическую\s+)?трубку\s+(\d+)\s*мм', name, re.IGNORECASE)
    if found:
        return 'Присоединение', f'{found.group(1)} мм'

    found = re.search(rf'D\s*=\s*(\d+(?:[.,]\d+)?\s*{X}\s*\d+(?:[.,]\d+)?)\s*мм', name)
    if found:
        return 'Диаметр', f'{latin(tidy(found.group(1)))} мм'

    found = re.search(r'D\s*=\s*(\d+(?:[.,]\d+)?)\s*мм', name)
    if found:
        return 'Диаметр', f'{tidy(found.group(1))} мм'

    return None, None


COLOURS = {
    'чёрн': 'чёрный', 'черн': 'чёрный', 'син': 'синий',
    'красн': 'красный', 'прозрачн': 'прозрачный', 'бел': 'белый',
}

MATERIALS = ('полиуретан', 'пластик', 'латунь', 'нержавеющая сталь',
             'алюминий', 'сталь')


def colour(name):
    found = re.search(r'цвет\s*:\s*([^\s,;]+)', name, re.IGNORECASE)
    if found:
        return tidy(found.group(1)).lower()
    for stem, value in COLOURS.items():
        if re.search(rf'{stem}(ый|ий|ая|яя|ое|ее|ые|ие)\b', name, re.IGNORECASE):
            return value
    return None


def plain_material(name):
    for value in MATERIALS:
        if re.search(rf'(?<![\w-]){value}', name, re.IGNORECASE):
            return value
    return None


def specifications(name, weight=None):
    """Return the pairs a single product name yields, in reading order."""
    name = name or ''
    pairs = []

    def add(label, value):
        if label and value:
            pairs.append((label, str(value)[:255]))

    add(*size(name))
    add('Резьба', thread(name))

    series = re.search(r'серии\s+(\d+)', name, re.IGNORECASE)
    add('Серия', series.group(1) if series else None)

    stroke = re.search(r'S\s*=\s*(\d+(?:[.,]\d+)?)\s*мм', name)
    add('Ход штока', f'{stroke.group(1)} мм' if stroke else None)

    # In a sensor's name the bare millimetres before the frequency are how far
    # it reaches: «6-30В, 2мм, 1500Гц, IP67»
    if re.search(r'датчик', name, re.IGNORECASE):
        reach = re.search(r'(\d+(?:[.,]\d+)?)\s*мм\s*,\s*\d+\s*Гц', name, re.IGNORECASE)
        add('Расстояние срабатывания', f'{reach.group(1)} мм' if reach else None)

    add('Температура', temperature(name))
    add('Давление', pressure(name))
    add('Давление на входе', pressure(name, prefix='вх'))
    add('Давление на выходе', pressure(name, prefix='вых'))

    flow = re.search(r'(\d+)\s*норм\.?\s*л/мин', name, re.IGNORECASE)
    add('Расход', f'{flow.group(1)} норм. л/мин' if flow else None)

    add('Напряжение', voltage(name))

    current = re.search(r'(\d+(?:[.,]\d+)?)\s*А(?![\w])', name)
    add('Ток', f'{current.group(1)} А' if current else None)

    frequency = re.search(r'(\d+(?:[-–]\d+)?)\s*Гц', name, re.IGNORECASE)
    add('Частота', f'{frequency.group(1)} Гц' if frequency else None)

    protection = re.search(r'\bIP\s*\d{2}\b', name, re.IGNORECASE)
    add('Степень защиты', tidy(protection.group(0)).upper().replace(' ', '')
        if protection else None)

    body = material(name, 'корпус')
    add('Материал корпуса', body)
    add('Материал цанги', material(name, 'цанга'))
    add('Материал уплотнений', material(name, 'уплотнения'))
    if not body:
        add('Материал', plain_material(name))

    add('Цвет', colour(name))

    if weight not in (None, '', 0):
        add('Вес', f'{weight} г')

    # One label per product, first mention wins
    seen, unique = set(), []
    for label, value in pairs:
        if label not in seen:
            seen.add(label)
            unique.append((label, value))
    return unique


def as_text(name, weight=None):
    return '\n'.join(f'{label}: {value}' for label, value in specifications(name, weight))


def add_column(source, target, name_col=1, weight_col=8):
    """Copy a catalogue file, adding a filled «Характеристики» column."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = openpyxl.load_workbook(source)
    sheet = book.active

    # Pictures are anchored to the photo column, so the new column goes at the
    # end rather than in the middle — the importer finds columns by their header
    column = sheet.max_column + 1
    header = sheet.cell(row=1, column=column, value='Характеристики')

    sample = sheet.cell(row=1, column=1)
    header.font = Font(bold=sample.font.bold, color=sample.font.color, size=sample.font.size)
    if sample.fill and sample.fill.fill_type:
        header.fill = PatternFill(start_color=sample.fill.start_color,
                                  end_color=sample.fill.end_color, fill_type='solid')
    header.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    header.border = Border(*(Side(style='thin'),) * 4)

    filled = 0
    for row in range(2, sheet.max_row + 1):
        text = as_text(sheet.cell(row=row, column=name_col).value,
                       sheet.cell(row=row, column=weight_col).value)
        cell = sheet.cell(row=row, column=column, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        filled += bool(text)

    sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 46
    book.save(target)
    return sheet.max_row - 1, filled


if __name__ == '__main__':
    import sys

    if len(sys.argv) != 3:
        sys.exit('Использование: python3 scripts/extract_specifications.py исходный.xlsx новый.xlsx')

    total, filled = add_column(sys.argv[1], sys.argv[2])
    print(f'{sys.argv[2]}: товаров {total}, характеристики заполнены у {filled}')
